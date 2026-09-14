import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_header(doc, exam_info, is_teacher=False):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Left column width 40%, Right 60%
    table.columns[0].width = Inches(3.0)
    table.columns[1].width = Inches(4.2)
    
    left_cell = table.cell(0, 0)
    right_cell = table.cell(0, 1)
    
    left_p = left_cell.paragraphs[0]
    left_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run1 = left_p.add_run(exam_info.get("header_left", "SỞ GIÁO DỤC VÀ ĐÀO TẠO\nTHÀNH PHỐ HẢI PHÒNG") + "\n")
    run1.font.name = "Times New Roman"
    run1.font.size = Pt(11)
    run1.font.bold = True
    
    run_sub = left_p.add_run("ĐỀ CHÍNH THỨC\n(Đề thi gồm 48 câu, 8 trang)")
    run_sub.font.name = "Times New Roman"
    run_sub.font.size = Pt(10)
    run_sub.font.italic = True
    
    right_p = right_cell.paragraphs[0]
    right_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = right_p.add_run(exam_info.get("exam_title", "KỲ THI CHỌN HỌC SINH GIỎI CẤP THÀNH PHỐ") + "\n")
    title_run.font.name = "Times New Roman"
    title_run.font.size = Pt(12)
    title_run.font.bold = True
    
    if is_teacher:
        t_tag = right_p.add_run("HƯỚNG DẪN CHẤM VÀ ĐÁP ÁN CHI TIẾT\n")
        t_tag.font.name = "Times New Roman"
        t_tag.font.size = Pt(11)
        t_tag.font.bold = True
        t_tag.font.color.rgb = RGBColor(180, 0, 0)
        
    sub_run = right_p.add_run(f"NĂM HỌC: {exam_info.get('school_year', '2026 - 2027')}\n")
    sub_run.font.name = "Times New Roman"
    sub_run.font.size = Pt(11)
    sub_run.font.bold = True
    
    subj_run = right_p.add_run(f"MÔN THI: {exam_info.get('subject', 'GIÁO DỤC KINH TẾ VÀ PHÁP LUẬT')}\n")
    subj_run.font.name = "Times New Roman"
    subj_run.font.size = Pt(11)
    subj_run.font.bold = True
    
    time_run = right_p.add_run(f"Thời gian làm bài: {exam_info.get('duration', '90 phút')} (không kể thời gian phát đề)\n")
    time_run.font.name = "Times New Roman"
    time_run.font.size = Pt(10)
    time_run.font.italic = True
    
    code_run = right_p.add_run(f"MÃ ĐỀ THI: {exam_info.get('code', '097')}")
    code_run.font.name = "Times New Roman"
    code_run.font.size = Pt(11)
    code_run.font.bold = True
    
    # Candidate line
    doc.add_paragraph()
    cand_p = doc.add_paragraph()
    cand_run = cand_p.add_run("Họ, tên thí sinh: .............................................................................. Số báo danh: ............................. Phòng thi: .............")
    cand_run.font.name = "Times New Roman"
    cand_run.font.size = Pt(11)
    cand_run.font.italic = True

def export_student_docx(exam, file_path):
    doc = Document()
    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    exam_info = exam.get("info", {})
    create_header(doc, exam_info, is_teacher=False)
    
    # PHẦN I
    p1_head = doc.add_paragraph()
    p1_head.paragraph_format.space_before = Pt(12)
    p1_head.paragraph_format.space_after = Pt(6)
    r = p1_head.add_run("PHẦN I. Câu trắc nghiệm nhiều phương án lựa chọn (6,0 điểm)\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r2 = p1_head.add_run("Thí sinh trả lời từ câu 1 đến câu 40. Mỗi câu hỏi thí sinh chỉ chọn một phương án.")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    r2.font.italic = True
    
    for q in exam.get("part1", []):
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(4)
        qp.paragraph_format.space_after = Pt(2)
        
        q_num = q.get("exam_number", 1)
        q_stem = q.get("stem", "")
        
        rq1 = qp.add_run(f"Câu {q_num}: ")
        rq1.font.name = "Times New Roman"
        rq1.font.size = Pt(11)
        rq1.font.bold = True
        
        rq2 = qp.add_run(q_stem)
        rq2.font.name = "Times New Roman"
        rq2.font.size = Pt(11)
        
        opts = q.get("options", {})
        for opt_key in ["A", "B", "C", "D"]:
            if opt_key in opts and opts[opt_key]:
                op = doc.add_paragraph()
                op.paragraph_format.left_indent = Inches(0.3)
                op.paragraph_format.space_after = Pt(1)
                ro1 = op.add_run(f"{opt_key}. ")
                ro1.font.name = "Times New Roman"
                ro1.font.size = Pt(11)
                ro1.font.bold = True
                ro2 = op.add_run(opts[opt_key])
                ro2.font.name = "Times New Roman"
                ro2.font.size = Pt(11)
                
    # PHẦN II
    p2_head = doc.add_paragraph()
    p2_head.paragraph_format.space_before = Pt(16)
    p2_head.paragraph_format.space_after = Pt(6)
    r_p2 = p2_head.add_run("PHẦN II. Câu trắc nghiệm đúng sai (4,0 điểm)\n")
    r_p2.font.name = "Times New Roman"
    r_p2.font.size = Pt(12)
    r_p2.font.bold = True
    r2_p2 = p2_head.add_run("Thí sinh trả lời từ câu 1 đến câu 8. Trong mỗi ý a), b), c), d) ở mỗi câu, thí sinh chọn đúng hoặc sai.")
    r2_p2.font.name = "Times New Roman"
    r2_p2.font.size = Pt(11)
    r2_p2.font.italic = True
    
    for q in exam.get("part2", []):
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(6)
        qp.paragraph_format.space_after = Pt(2)
        
        q_num = q.get("exam_number", 1)
        q_stem = q.get("stem", "")
        
        rq1 = qp.add_run(f"Câu {q_num}: ")
        rq1.font.name = "Times New Roman"
        rq1.font.size = Pt(11)
        rq1.font.bold = True
        
        rq2 = qp.add_run(q_stem)
        rq2.font.name = "Times New Roman"
        rq2.font.size = Pt(11)
        
        for stmt in q.get("statements", []):
            sp = doc.add_paragraph()
            sp.paragraph_format.left_indent = Inches(0.3)
            sp.paragraph_format.space_after = Pt(1)
            lbl = stmt.get("label", "a")
            txt = stmt.get("text", "")
            
            rs1 = sp.add_run(f"{lbl}) ")
            rs1.font.name = "Times New Roman"
            rs1.font.size = Pt(11)
            rs1.font.bold = True
            
            rs2 = sp.add_run(txt)
            rs2.font.name = "Times New Roman"
            rs2.font.size = Pt(11)
            
    # End note
    end_p = doc.add_paragraph()
    end_p.paragraph_format.space_before = Pt(18)
    end_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rend = end_p.add_run("---------- HẾT ----------\nCán bộ coi thi không giải thích gì thêm.")
    rend.font.name = "Times New Roman"
    rend.font.size = Pt(11)
    rend.font.italic = True
    
    doc.save(file_path)
    return file_path

def export_teacher_docx(exam, file_path):
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    exam_info = exam.get("info", {})
    create_header(doc, exam_info, is_teacher=True)
    
    # 1. Answer Key Summary Table for Part I
    p1_sum = doc.add_paragraph()
    p1_sum.paragraph_format.space_before = Pt(14)
    r1 = p1_sum.add_run("I. BẢNG ĐÁP ÁN PHẦN I (TRẮC NGHIỆM NHIỀU LỰA CHỌN)")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    p1_list = exam.get("part1", [])
    cols = 10
    rows = (len(p1_list) + cols - 1) // cols
    
    tbl1 = doc.add_table(rows=rows * 2, cols=cols)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    for r_idx in range(rows):
        # Header row (Câu 1, Câu 2...)
        for c_idx in range(cols):
            q_idx = r_idx * cols + c_idx
            cell_q = tbl1.cell(r_idx * 2, c_idx)
            set_cell_background(cell_q, "E8EEF5")
            p_q = cell_q.paragraphs[0]
            p_q.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            cell_a = tbl1.cell(r_idx * 2 + 1, c_idx)
            p_a = cell_a.paragraphs[0]
            p_a.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            if q_idx < len(p1_list):
                q = p1_list[q_idx]
                rq = p_q.add_run(f"{q.get('exam_number', q_idx + 1)}")
                rq.font.name = "Times New Roman"
                rq.font.size = Pt(10)
                rq.font.bold = True
                
                ans = q.get("answer", "A")
                ra = p_a.add_run(ans)
                ra.font.name = "Times New Roman"
                ra.font.size = Pt(10)
                ra.font.bold = True
                ra.font.color.rgb = RGBColor(180, 0, 0)
                
    # 2. Answer Key Summary Table for Part II
    doc.add_paragraph()
    p2_sum = doc.add_paragraph()
    r2 = p2_sum.add_run("II. BẢNG ĐÁP ÁN PHẦN II (TRẮC NGHIỆM ĐÚNG / SAI)")
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(12)
    r2.font.bold = True
    
    p2_list = exam.get("part2", [])
    tbl2 = doc.add_table(rows=len(p2_list) + 1, cols=5)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Câu hỏi", "Lệnh hỏi a", "Lệnh hỏi b", "Lệnh hỏi c", "Lệnh hỏi d"]
    for c_idx, h in enumerate(headers):
        c = tbl2.cell(0, c_idx)
        set_cell_background(c, "E8EEF5")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rh = p.add_run(h)
        rh.font.name = "Times New Roman"
        rh.font.size = Pt(10)
        rh.font.bold = True
        
    for idx, q in enumerate(p2_list, 1):
        c_q = tbl2.cell(idx, 0)
        pq = c_q.paragraphs[0]
        pq.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rq = pq.add_run(f"Câu {q.get('exam_number', idx)}")
        rq.font.name = "Times New Roman"
        rq.font.size = Pt(10)
        rq.font.bold = True
        
        stmts = q.get("statements", [])
        for s_idx, lbl in enumerate(["a", "b", "c", "d"], 1):
            cell_ans = tbl2.cell(idx, s_idx)
            pa = cell_ans.paragraphs[0]
            pa.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Find statement
            matching = [s for s in stmts if s.get("label") == lbl]
            ans_str = "Đ" if matching and matching[0].get("answer") is True else "S"
            ra = pa.add_run(ans_str)
            ra.font.name = "Times New Roman"
            ra.font.size = Pt(10)
            ra.font.bold = True
            if ans_str == "Đ":
                ra.font.color.rgb = RGBColor(0, 120, 0)
            else:
                ra.font.color.rgb = RGBColor(180, 0, 0)
                
    # 3. Detailed Explanation & Legal Basis
    doc.add_paragraph()
    p_exp = doc.add_paragraph()
    r_exp = p_exp.add_run("III. HƯỚNG DẪN GIẢI CHI TIẾT & CĂN CỨ PHÁP LÝ")
    r_exp.font.name = "Times New Roman"
    r_exp.font.size = Pt(12)
    r_exp.font.bold = True
    
    for q in p1_list:
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(4)
        q_num = q.get("exam_number", 1)
        rq = qp.add_run(f"Câu {q_num} (Đáp án {q.get('answer', 'A')} - Mức độ {q.get('level', '').upper()}): ")
        rq.font.name = "Times New Roman"
        rq.font.size = Pt(10)
        rq.font.bold = True
        
        exp_text = q.get("explanation", "Căn cứ quy định pháp luật hiện hành và nội dung chuyên đề.")
        re = qp.add_run(exp_text)
        re.font.name = "Times New Roman"
        re.font.size = Pt(10)
        
    for q in p2_list:
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(6)
        q_num = q.get("exam_number", 1)
        rq = qp.add_run(f"Câu {q_num} (Phần Đúng/Sai):\n")
        rq.font.name = "Times New Roman"
        rq.font.size = Pt(10)
        rq.font.bold = True
        
        for s in q.get("statements", []):
            sp = doc.add_paragraph()
            sp.paragraph_format.left_indent = Inches(0.3)
            lbl = s.get("label", "a")
            ans_tag = "ĐÚNG" if s.get("answer") else "SAI"
            rs1 = sp.add_run(f"Ý {lbl}) [{ans_tag}]: ")
            rs1.font.name = "Times New Roman"
            rs1.font.size = Pt(10)
            rs1.font.bold = True
            
            rs2 = sp.add_run(s.get("explanation", ""))
            rs2.font.name = "Times New Roman"
            rs2.font.size = Pt(10)
            
    doc.save(file_path)
    return file_path

def export_matrix_docx(exam, file_path):
    doc = Document()
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title_p.add_run("MA TRẬN VÀ BẢN ĐẶC TẢ ĐỀ THI HỌC SINH GIỎI\nMÔN: GIÁO DỤC KINH TẾ VÀ PHÁP LUẬT")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True
    
    doc.add_paragraph()
    # Stats table
    stats = exam.get("stats", {})
    p_info = doc.add_paragraph()
    r_info = p_info.add_run(f"Tổng số câu hỏi: {stats.get('total_questions', 48)} câu | Phần I: {stats.get('part1_total', 40)} câu TN | Phần II: {stats.get('part2_total', 8)} câu Đúng/Sai (32 ý)")
    r_info.font.name = "Times New Roman"
    r_info.font.size = Pt(11)
    r_info.font.italic = True
    
    # Add matrix table
    tbl = doc.add_table(rows=5, cols=8)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Mạch kiến thức", "Nội dung", "Phần I (Biết)", "Phần I (Hiểu)", "Phần I (Vận dụng)", "Phần II (Biết)", "Phần II (Hiểu)", "Phần II (Vận dụng)"]
    for c_idx, h in enumerate(headers):
        c = tbl.cell(0, c_idx)
        set_cell_background(c, "E8EEF5")
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rh = p.add_run(h)
        rh.font.name = "Times New Roman"
        rh.font.size = Pt(10)
        rh.font.bold = True
        
    rows_data = [
        ("Giáo dục Pháp luật 10", "Pháp luật nước CHXHCN Việt Nam", "2", "3", "1", "2", "4", "2"),
        ("Giáo dục Pháp luật 11", "Bình đẳng, Dân chủ và Tự do", "3", "7", "7", "3", "5", "4"),
        ("Giáo dục Kinh tế 12", "Tăng trưởng KT, Hội nhập, BHXH", "2", "6", "9", "2", "4", "6"),
        ("TỔNG CỘNG", "Toàn đề (48 câu)", "7", "16", "17", "7", "13", "12")
    ]
    for r_idx, r_vals in enumerate(rows_data, 1):
        for c_idx, val in enumerate(r_vals):
            c = tbl.cell(r_idx, c_idx)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx >= 2 else WD_ALIGN_PARAGRAPH.LEFT
            r_run = p.add_run(val)
            r_run.font.name = "Times New Roman"
            r_run.font.size = Pt(10)
            if r_idx == 4 or c_idx == 0:
                r_run.font.bold = True
                
    doc.save(file_path)
    return file_path
